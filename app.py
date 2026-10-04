import io
import re

import streamlit as st
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

st.set_page_config(page_title="AI Generator", layout="wide")

st.markdown("""
    <style>
    div.stButton > button:first-child {
        background-color: #a8dfb5 !important;
        color: #1e4620 !important;
        border: 1px solid #8ecda0 !important;
        font-weight: bold !important;
        font-size: 18px !important;
        padding: 10px 24px !important;
        border-radius: 8px !important;
        transition: all 0.3s ease !important;
    }
    div.stButton > button:first-child:hover {
        background-color: #92cfa0 !important;
        border-color: #79be8c !important;
        box-shadow: 0 4px 12px rgba(0,0,0,0.08) !important;
    }
    </style>
    """, unsafe_allow_html=True)

STYLES = {
    "Минималистичный светлый": {"bg": "FFFFFF", "fg": "222222", "accent": "2E7D32"},
    "Строгий корпоративный": {"bg": "F4F6F8", "fg": "1F2933", "accent": "1F4E79"},
    "Яркий креативный": {"bg": "FFF4E0", "fg": "3A2A1A", "accent": "E4572E"},
    "Технологичный темный футуризм": {"bg": "0D1117", "fg": "E6EDF3", "accent": "00E5FF"},
}

# ---------- Состояние ----------
st.session_state.setdefault("slider_input", 5)
st.session_state.setdefault("num_input", 5)
st.session_state.setdefault("user_text", "")
st.session_state.setdefault("uploader_key", 0)
st.session_state.setdefault("files_sig", ())
st.session_state.setdefault("result", None)


def sync_from_slider():
    st.session_state.num_input = st.session_state.slider_input


def sync_from_num():
    st.session_state.slider_input = st.session_state.num_input


def reset_form():
    st.session_state.slider_input = 5
    st.session_state.num_input = 5
    st.session_state.user_text = ""
    st.session_state.files_sig = ()
    st.session_state.result = None
    st.session_state.uploader_key += 1


# ---------- Логика генерации ----------
def read_text_file(f):
    raw = f.getvalue()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1251", errors="replace")


def split_text(text, groups):
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paras:
        return []
    groups = max(1, min(groups, len(paras)))
    size = -(-len(paras) // groups)
    return [paras[i:i + size] for i in range(0, len(paras), size)]


def clean(line, limit):
    line = line.strip().lstrip("#-*• ").strip()
    return line if len(line) <= limit else line[:limit - 1] + "…"


def chunk_to_slide(chunk):
    first_lines = [l for l in chunk[0].splitlines() if l.strip()]
    title = clean(first_lines[0], 70)
    body = [clean(l, 220) for l in first_lines[1:]]
    for para in chunk[1:]:
        body += [clean(l, 220) for l in para.splitlines() if l.strip()]
    return title, body[:8]


def clear_slides(prs):
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def body_placeholder(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            return ph
    return None


def set_title(slide, text):
    if slide.shapes.title is not None:
        slide.shapes.title.text = text
    else:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8.6), Inches(1))
        box.text_frame.text = text


def set_body(slide, lines, size=18):
    ph = body_placeholder(slide)
    if ph is None:
        ph = slide.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(8.2), Inches(5))
    tf = ph.text_frame
    tf.word_wrap = True
    tf.text = lines[0] if lines else ""
    for line in lines[1:]:
        tf.add_paragraph().text = line
    for p in tf.paragraphs:
        p.font.size = Pt(size)


def apply_style(slide, palette):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(palette["bg"])
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        is_title = shape == slide.shapes.title
        color = palette["accent"] if is_title else palette["fg"]
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor.from_string(color)
                if is_title:
                    r.font.bold = True


def build_presentation(text, total, style_name, template_bytes):
    if template_bytes:
        prs = Presentation(io.BytesIO(template_bytes))
        clear_slides(prs)
    else:
        prs = Presentation()

    layouts = prs.slide_layouts
    title_layout = layouts[0]
    content_layout = layouts[1] if len(layouts) > 1 else layouts[0]

    chunks = split_text(text, max(total - 1, 1))
    deck_title = clean(chunks[0][0].splitlines()[0], 70) if chunks else "ИИ Генерация"

    slide = prs.slides.add_slide(title_layout)
    set_title(slide, deck_title)
    set_body(slide, [f"Стиль: {style_name}"], size=20)
    slides = [slide]

    for chunk in chunks[: max(total - 1, 0)]:
        title, body = chunk_to_slide(chunk)
        s = prs.slides.add_slide(content_layout)
        set_title(s, title)
        set_body(s, body)
        slides.append(s)

    while len(slides) < total:
        s = prs.slides.add_slide(content_layout)
        set_title(s, "Дополните слайд")
        set_body(s, ["Добавьте больше текста, чтобы заполнить этот слайд."])
        slides.append(s)

    if not template_bytes:
        for s in slides:
            apply_style(s, STYLES[style_name])

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ---------- Боковая панель ----------
with st.sidebar:
    st.markdown("## 📥 Загрузка материалов")
    st.caption("Можно выбрать несколько файлов сразу. Первый PPTX используется как шаблон.")
    uploaded_files = st.file_uploader(
        "Выберите файлы PPTX, TXT или MD",
        type=["pptx", "txt", "md"],
        accept_multiple_files=True,
        key=f"file_uploader_{st.session_state.uploader_key}",
    )
    st.divider()
    st.button("❌ Очистить форму", on_click=reset_form, use_container_width=True)

# ---------- Чтение файлов ----------
template_bytes = None
texts = []
for f in uploaded_files or []:
    ext = f.name.rsplit(".", 1)[-1].lower()
    if ext in ("txt", "md"):
        texts.append(f"# {f.name}\n{read_text_file(f)}")
    elif ext == "pptx" and template_bytes is None:
        template_bytes = f.getvalue()

sig = tuple((f.name, f.size) for f in (uploaded_files or []))
if sig != st.session_state.files_sig:
    prev = st.session_state.files_sig
    st.session_state.files_sig = sig
    if texts or prev:
        st.session_state.user_text = "\n\n".join(texts)

# ---------- Основная область ----------
st.markdown("# 🧠 ИИ-Ассистент презентаций")
st.info(
    "Загрузите текстовые документы и при желании шаблон PPTX в боковом меню. "
    "Тексты автоматически объединятся и появятся в поле ниже."
)
st.divider()

user_text = st.text_area(
    "Содержание презентации:",
    key="user_text",
    placeholder="Введите структуру вручную или загрузите файлы слева. Абзацы разделяйте пустой строкой.",
    height=150,
)

col1, col2 = st.columns(2)
with col1:
    st.markdown("**Настройка объема слайдов**")
    st.slider(
        "Выберите количество на шкале:",
        min_value=1, max_value=50,
        key="slider_input", on_change=sync_from_slider,
    )
    st.number_input(
        "Или введите точное число (от 1 до 50):",
        min_value=1, max_value=50,
        key="num_input", on_change=sync_from_num,
    )
with col2:
    st.markdown("**Визуальное оформление**")
    style_option = st.selectbox(
        "Выберите подходящий визуальный стиль:", list(STYLES.keys())
    )
    st.caption("Стиль задаёт цвета слайдов. Если загружен шаблон, действуют его цвета.")

st.divider()

if st.button("✨ Запустить генерацию презентации", use_container_width=True):
    if user_text.strip() or template_bytes:
        with st.spinner("Формирую слайды..."):
            total = int(st.session_state.slider_input)
            try:
                data = build_presentation(user_text, total, style_option, template_bytes)
                st.session_state.result = {
                    "data": data,
                    "name": f"ai_presentation_{total}_slides.pptx",
                }
            except Exception as e:
                st.session_state.result = None
                st.error(f"Не удалось создать презентацию: {e}")
    else:
        st.error("⚠️ Загрузите файлы в левой панели или введите текст вручную.")

# Блок результата вынесен из if-кнопки: после скачивания он не пропадает
if st.session_state.result:
    res = st.session_state.result
    st.success("🎉 Презентация сформирована.")
    st.markdown(f"📄 **Файл:** `{res['name']}`")
    st.download_button(
        label="📥 Скачать готовую презентацию (PPTX)",
        data=res["data"],
        file_name=res["name"],
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        use_container_width=True,
    )
